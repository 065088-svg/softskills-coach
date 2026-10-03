from docx import Document
from docx.shared import Pt
d = Document()
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)

def H(t, l=1): d.add_heading(t, l)
def P(t): d.add_paragraph(t)
def B(t): d.add_paragraph(t, style="List Bullet")
def Q(q, a):
    p = d.add_paragraph(); r = p.add_run(q); r.bold = True
    for para in (a if isinstance(a, list) else [a]):
        if para.startswith("- "): B(para[2:])
        else: P(para)

d.add_heading("SoftSkills Coach — Project Report", 0)
P("AI Application End Term Project (30 marks) · Prof. Vaibhav Jain\nUse case 15: Language / soft-skills practice bot (Chatbot)\nStudent: Lakshya · PGDM (BDA) Batch-06, FORE School of Management\nApp link: [PASTE YOUR STREAMLIT URL]   ·   Demo video: [PASTE LINK]")

H("About the product")
P("SoftSkills Coach lets a learner rehearse realistic workplace conversations (job interview, upset customer, workload negotiation with a manager, networking with a senior professional, giving peer feedback, a client call). Every turn is multiple choice: the learner picks one of four replies, instantly sees a rating (Best / Acceptable / Weak / Poor), why, and a rule of thumb. Each question tests one skill (Grammar, Tone, Clarity, Empathy, Assertiveness). Scores and per-skill accuracy are saved across sessions. Sample data: 6 scenarios × 5 questions = 30 hand-written questions (120 rated replies).")
P("Two question sources: (1) Verified — hand-written, fixed ratings, works offline; (2) Live AI — Gemini writes each next turn from the conversation so far, with every output validated before display. Gemini also writes a short coaching note at the end of a session.")

H("A. Business & Strategic Framing (SWOT)")
Q("A1. Problem, for whom, who pays?", [
 "Problem: professionals and students know what good communication looks like but rarely get safe, repeated, feedback-rich practice. Mock interviews and coaches are scarce and expensive; apps like flashcards do not rehearse real situations.",
 "End user: MBA/graduate students and early-career employees (especially non-native English speakers in business settings). Paying customer: in the B2C case the learner; in the B2B/B2B2C case, the institution — a business school placement cell, a corporate L&D/HR team, or a BPO/IT-services training department — which pays per seat."])
Q("A2. Strengths — what does it do better than a human or competing tool?", [
 "- Available 24×7, unlimited repetition, no embarrassment in front of a coach.",
 "- Instant, consistent feedback that explains why an answer is better (a human coach cannot give the same rating to the same answer every time; this bot can).",
 "- Measures progress by skill across sessions, so the learner sees where to focus.",
 "- Accuracy by design: in Verified mode the score never depends on an AI’s judgement; it comes from ratings written and reviewed by the author."])
Q("A3. Weaknesses — where does it break or give unreliable output? (example)", [
 "- Rating gaps are judgement: in the interview scenario, question 5 rates “No, I think you’ve covered everything” (Weak) above “When will I get promoted, and can I avoid late hours?” (Poor). Another coach could order two imperfect answers differently. Only the Best answer per question is meant to be unambiguous.",
 "- Live AI mode ratings are model judgement. The app checks structure (exactly four replies, ratings 3/2/1/0 once each, length limits) but cannot verify that the model’s grammar claim is true. This is why Verified is the default and why every question has a “Report this question” button that saves the item for human review.",
 "- Content is limited to 6 scenarios in Verified mode, in English only, with one cultural register (international business English).",
 "[ADD: a screenshot of any Live-mode question you personally find questionable during testing, if any.]"])
Q("A4. Opportunities — how would it scale?", [
 "- More users: move storage from a local JSON file to a database (Postgres/Firestore) and add logins.",
 "- More languages: the scenario schema is language-neutral; Gemini can draft Hindi/Spanish/German scenarios which a native reviewer approves into the Verified bank.",
 "- More use cases: sales calls, medical-student patient communication, customer-support onboarding, campus placement preparation; add voice (speech-to-text) later.",
 "- Admin dashboard for institutions showing cohort-level skill gaps."])
Q("A5. Threats — what could kill it?", [
 "- Large incumbents adding the same feature (e.g., language apps with AI role-play).",
 "- API price or free-tier changes, and model deprecation. This already happened during the build: current Google documentation states Gemini 2.5 models are available only to past users, so the app does not hard-code one model; it tries a list, then auto-discovers Flash models available to the key.",
 "- Data-protection regulation (India’s DPDP Act, GDPR) if real learner data is stored at scale.",
 "- Trust: one visibly wrong grammar rating can damage credibility, which is why ratings are hand-verified and reportable."])
Q("A6. Two real competitors and the difference", [
 "- Duolingo (Max tier offers AI role-play): focused on language learning for general life; gamified. SoftSkills Coach targets workplace soft skills with explicit tone/assertiveness/empathy scoring and an MCQ format that gives deterministic marks.",
 "- Yoodli (AI communication/roleplay coaching): speech-based, aimed at sales and leadership teams, priced for enterprises. SoftSkills Coach is lightweight, free-tier-friendly, text/MCQ-based and aimed at students and early-career learners.",
 "(Competitor features are summarised from their public marketing; verify the latest details before presenting.)"])
Q("A7. Monetisation / adoption path", [
 "- Phase 1: free for students, adopted through placement cells (pilot at the college).",
 "- Phase 2: B2B2C — per-seat annual licence to institutions and corporate L&D teams, with analytics dashboard and custom scenarios (e.g., a company’s own customer-complaint cases).",
 "- Phase 3: premium for individuals (more scenarios, voice practice, certificates). Cost structure is low: each Live session uses about six small Flash-class calls."])

H("B. AI / Technical Understanding")
Q("B1. Which model/API and why?", [
 "Google Gemini via the google-genai SDK, using Flash / Flash-Lite class models. Reasons: the project brief asks for a Gemini free key; Flash-Lite is the lowest-latency, lowest-cost tier and our task (short structured JSON, short coaching note) does not need a large model; Gemini supports JSON-schema structured output, which lets the app enforce the shape of every question. Alternatives considered: larger Pro-class models (higher cost/latency, no benefit for 4-option questions) and OpenAI/Claude APIs (no free key requirement fit). Model choice is configurable (GEMINI_MODEL) and the client falls back across models automatically."])
Q("B2. Prompt / system design and guardrails", [
 "- System prompt defines the role (content engine for a practice app), forces JSON-only output, and fixes the rubric: 3 = Best … 0 = Poor, each used exactly once per question.",
 "- Constraints: partner line ≤ 45 words; options 1–3 sentences; one-sentence feedback naming the grammar rule; “if unsure a form is an error, do not use it as one”; do not make the best option the longest; British spelling; no real people, brands or personal data.",
 "- Conversation history (partner line, learner’s reply, reply quality) is passed as JSON data, and the prompt states that everything in the user message is data, never instructions.",
 "- Response schema (structured output) plus code validation: exactly 4 options, integer ratings 0–3 used once each, no duplicates, length limits, control characters stripped. One automatic retry that tells the model what was wrong; then fall back to Verified questions.",
 "- Debrief prompt: uses only the supplied results; no markdown; no personality comments; output cleaned and length-capped."])
Q("B3. Out-of-scope input — demonstrate", [
 "The app has no free-text chat box, so a user cannot send out-of-scope text to the model at all; every input is one of four buttons. The only typed field is the display name, validated by a strict pattern (2–20 letters/numbers/space . _ -). Test cases in the test-suite show “<script>alert(1)</script>”, “../../etc” and “Ignore all previous instructions now” are rejected. Model output is also treated as untrusted: it is validated and sanitised before display.",
 "[ADD: screenshot of the name field rejecting “<script>”.]"])
Q("B4. Data privacy", [
 "Yes, third-party processing exists and is disclosed in the sidebar (“Privacy & AI disclosure”) and in the README. Verified mode: nothing is sent to Google until the end-of-session coach note, when scores and the chosen/best reply texts are sent (only if a key is set). Live mode: scenario text and the learner’s chosen replies are sent to the Gemini API. The learner’s name is never sent (verified by an automated test that inspects every outbound request). Progress is stored locally in data/progress.json. The free Gemini tier may use submitted content to improve Google products under Google’s terms, so no sensitive personal data should be entered — and the app never asks for any."])
Q("B5. Failure mode — API down or garbage output", [
 "Every failure ends in a playable session. Tested against a local fake Gemini server using the real SDK: wrong key → clear message; quota (429) or model not found (404) → next model, then auto-discovery; server error → one retry then next model; network down → message; invalid or empty output → one corrective retry, then fallback to Verified questions with an on-screen notice; failure on the final coach note → the deterministic summary is still shown. Progress is saved after every answer, so a crash or refresh offers “Resume”."])
Q("B6. RAG / fine-tuning?", "Not used. Accuracy comes from a hand-written, rubric-based question bank (Verified mode) with automated structural audits (4 options, ratings 3/2/1/0 once each, no duplicate texts, shuffled positions) rather than retrieval or fine-tuning.")

H("C. Critical Thinking / Honest Capability Assessment")
Q("C1. Example where the bot was confidently wrong or misleading", [
 "Honest answer from my own build: my first version of the question bank was misleading in a different way — the Best reply was the longest option in 28 of 30 questions, so a learner could score highly just by choosing the longest reply without learning anything, while the app confidently showed “Best”. An automated audit caught this; I rewrote ~90 options so the Best reply’s length rank is spread (1st: 9, 2nd: 11, 3rd: 5, 4th: 5 questions) and added a permanent test.",
 "[ADD if you observe one: a screenshot of a Live-mode question where Gemini’s rating or grammar explanation looked wrong.]"])
Q("C2. What would you NOT trust it to do unsupervised?", "Formal assessment of candidates (hiring or grading decisions), certifying someone’s language proficiency, or judging culture-specific etiquette. Ratings reflect one author’s view of international business English. Live-mode ratings should always be treated as suggestions.")
Q("C3. Who is accountable if a user acts on bad output?", "The product owner (me) is accountable for the correctness of the Verified bank and for disclosing limitations; the model provider is accountable for the model’s general behaviour under its terms; the user remains responsible for real-world decisions. The app reduces risk by labelling AI-generated content, providing a “Report this question” button, and positioning itself as practice, not certification.")
Q("C4. Biggest limitation of the underlying model designed around", "LLMs can state wrong grammar rules fluently and are inconsistent when asked to grade. Design response: do not let the model decide scores in the default mode; in Live mode constrain the output with a schema and rubric, validate it, and keep a human-reviewed fallback.")

H("D. Execution")
Q("D1. Edge case tested and what changed", [
 "- Length bias (above): test found the best reply was longest in 28/30 questions → rewrote options and added a distribution test.",
 "- Model availability: Google documentation shows older Gemini 2.5 models are restricted to past users, so a hard-coded model would fail for a new key → added model fallback and auto-discovery plus check_key.py.",
 "- Refresh/double-click: a double submit is ignored (idempotent scoring) and unfinished sessions are resumed after a refresh.",
 "- Corrupted progress file: quarantined to a .corrupt file and the app starts fresh instead of crashing.",
 "User ease of use: 5 steps in the README, a single .env file for the key, a sidebar fallback box for the key, and a one-command key test (python check_key.py)."])

H("F. Chatbot-Specific — Conversation Design & Flow")
Q("F1. Multi-turn memory", "Yes. In Live mode the full history of the session (what each partner said, what the learner chose and how good it was) is passed to Gemini on every turn so the partner reacts to earlier choices. In Verified mode the partner lines are written to make sense after any previous reply, and a rating-based reaction line (e.g., “Ms. Rao raises an eyebrow slightly.”) gives feedback continuity; the transcript of earlier turns stays on screen.")
Q("F2. Fallback when intent is not understood", "There is no intent to misunderstand because input is multiple choice. If the model fails to produce a valid question, the bot does not guess: it silently switches to a verified question and shows a notice. If the user presses Submit without choosing, it asks them to select a reply.")
Q("F3. Off-topic or adversarial message", "Not possible through the interface (no free text). See B3: the only text field rejects markup and path strings; model output and imported progress files are validated; an imported file with tampered scores is rejected record by record. [ADD: screenshot of the rejection message.]")
Q("F4. How does the user know it is an AI, and is there a hand-off?", "A banner says “AI-powered practice tool”, every partner message is labelled “(simulated)”, the sidebar has a privacy & AI disclosure, and AI-written coach notes are boxed and labelled “written by Gemini … may contain mistakes”. Human hand-off: the “Report this question” button saves the item to a review queue for the author.")
Q("F5. Persona, tone and why", "Partners are realistic but professional characters (an HR manager, an upset customer, a manager…). The coaching voice is encouraging, specific and brief: feedback is one sentence plus a rule of thumb. This suits adult learners who want quick, actionable correction rather than long lectures.")
Q("F6. Recovery from vague or incomplete answers", "Because replies are chosen from four options, vagueness is modelled instead of left to chance: each question includes a vague option, which is rated lower with an explanation of what is missing (e.g., “‘Soon’ gives no deadline”). The review screen shows the Best reply next to what the learner chose.")
Q("F7. Consistency for the same question phrased differently", "Verified mode is deterministic: the same option always gets the same rating and feedback (option order is shuffled, ratings stay attached — tested). Since users cannot rephrase questions, consistency across phrasings does not apply; in Live mode the same scenario may produce different questions each time, which is expected variety, but ratings are produced under the same fixed rubric.")

H("Appendix: Testing summary")
P("147 automated tests pass: question-bank audit (structure, unique texts, shuffling, length-bias distribution), scoring (perfect/zero/mixed hand-calculated), input validation, storage (corruption, import/export, isolation), Gemini integration through the real SDK against a local fake server (model fallback, auth, quota, retries, invalid output), and end-to-end UI flows (login, full session, resume after refresh, live mode, API failure mid-session, privacy check).")
P("Limitation to state openly: the live Gemini API could not be called from the build environment, so Live mode was verified against a fake server that mimics the API; run python check_key.py on your machine to confirm your key and model.")
d.save("docs/SoftSkills_Coach_Project_Report.docx")
print("saved")
